from pathlib import Path
from datetime import datetime

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

import ticket_service
from excel_handler import (
    RESULT_HEADERS,
    append_result,
    find_first_recorded_ticket,
    get_group_names,
    get_students,
)
from ticket_parser import TicketParseError, parse_tickets
from ticket_service import TicketSaveError, TicketService


def create_document(
    file_path: Path, tickets: dict[int, list[str]]
) -> None:
    document = Document()
    for ticket_number, questions in tickets.items():
        document.add_paragraph(f"Билет {ticket_number}")
        for question_number, question in enumerate(questions, start=1):
            document.add_paragraph(f"{question_number}. {question}")
    document.save(file_path)


def create_students_workbook(
    file_path: Path, groups: dict[str, list[tuple[str, str]]]
) -> None:
    workbook = Workbook()
    for index, (group_name, students) in enumerate(groups.items()):
        worksheet = (
            workbook.active
            if index == 0
            else workbook.create_sheet()
        )
        worksheet.title = group_name
        worksheet.append(["Фамилия", "Имя"])
        for last_name, first_name in students:
            worksheet.append([last_name, first_name])
    workbook.save(file_path)
    workbook.close()


@pytest.fixture
def valid_ticket_data() -> dict[int, list[str]]:
    return {
        2: ["Вопрос 2.1", "Вопрос 2.2", "Вопрос 2.3"],
        11: ["Вопрос 11.1", "Вопрос 11.2", "Вопрос 11.3"],
    }


def test_word_parser_reads_two_valid_tickets(
    tmp_path: Path, valid_ticket_data: dict[int, list[str]]
) -> None:
    document_path = tmp_path / "tickets.docx"
    create_document(document_path, valid_ticket_data)

    parsed_tickets = parse_tickets(document_path)

    assert parsed_tickets == valid_ticket_data
    assert all(len(questions) == 3 for questions in parsed_tickets.values())


def test_word_parser_skips_ticket_with_too_few_questions(
    tmp_path: Path, valid_ticket_data: dict[int, list[str]]
) -> None:
    document_path = tmp_path / "tickets.docx"
    document = Document()
    document.add_paragraph("Билет 5")
    document.add_paragraph("1. Недостаточно вопросов")
    document.add_paragraph("Билет 11")
    for number, question in enumerate(valid_ticket_data[11], start=1):
        document.add_paragraph(f"{number}. {question}")
    document.save(document_path)

    assert parse_tickets(document_path) == {11: valid_ticket_data[11]}


def test_student_group_lookup_returns_only_selected_group(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "students.xlsx"
    create_students_workbook(
        workbook_path,
        {
            "Группа А": [("Иванов", "Иван"), ("Петрова", "Анна")],
            "Группа Б": [("Сидоров", "Пётр")],
        },
    )

    assert get_group_names(workbook_path) == ["Группа А", "Группа Б"]
    assert get_students(workbook_path, "Группа Б") == [
        {"last_name": "Сидоров", "first_name": "Пётр"}
    ]


def test_empty_student_group_returns_empty_list(tmp_path: Path) -> None:
    workbook_path = tmp_path / "students.xlsx"
    create_students_workbook(workbook_path, {"Пустая группа": []})

    assert get_students(workbook_path, "Пустая группа") == []


def test_first_ticket_generation_is_saved_as_new_record(
    tmp_path: Path,
    valid_ticket_data: dict[int, list[str]],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    tickets_path = tmp_path / "tickets.docx"
    results_path = tmp_path / "results.xlsx"
    create_document(tickets_path, valid_ticket_data)
    monkeypatch.setattr(ticket_service.random, "choice", lambda numbers: 11)

    assignment = TicketService(tickets_path, results_path).generate_ticket(
        "Группа", "Иванов", "Иван"
    )

    assert assignment.ticket_number in valid_ticket_data
    assert assignment.is_repeat is False
    workbook = load_workbook(results_path)
    worksheet = workbook.active
    assert worksheet.max_row == 2
    assert worksheet["A2"].value == "Группа"
    assert worksheet["B2"].value == "Иванов"
    assert worksheet["C2"].value == "Иван"
    assert worksheet["D2"].value == assignment.ticket_number
    assert isinstance(worksheet["E2"].value, datetime)
    assert worksheet["E2"].number_format == "DD.MM.YYYY HH:MM"
    assert worksheet.column_dimensions["E"].width == 22
    assert worksheet.column_dimensions["B"].width >= len("Иванов") + 2
    assert worksheet["F2"].value == "нет"
    workbook.close()


def test_repeated_generation_reuses_ticket_and_appends_row(
    tmp_path: Path,
    valid_ticket_data: dict[int, list[str]],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    tickets_path = tmp_path / "tickets.docx"
    results_path = tmp_path / "results.xlsx"
    create_document(tickets_path, valid_ticket_data)
    monkeypatch.setattr(ticket_service.random, "choice", lambda numbers: 2)
    service = TicketService(tickets_path, results_path)

    first_assignment = service.generate_ticket("Группа", "Иванов", "Иван")
    second_assignment = service.generate_ticket("Группа", "Иванов", "Иван")

    assert second_assignment.ticket_number == first_assignment.ticket_number
    assert second_assignment.is_repeat is True
    workbook = load_workbook(results_path)
    assert workbook.active.max_row == 3
    assert workbook.active["D3"].value == first_assignment.ticket_number
    assert workbook.active["F3"].value == "да"
    workbook.close()


def test_journal_preserves_existing_rows_after_reopen(tmp_path: Path) -> None:
    results_path = tmp_path / "results.xlsx"
    for student_number in range(3):
        append_result(
            results_path,
            "Группа",
            f"Фамилия{student_number}",
            f"Имя{student_number}",
            student_number + 1,
        )

    workbook = load_workbook(results_path)
    original_rows = [
        tuple(cell.value for cell in row)
        for row in workbook.active.iter_rows(min_row=2, max_row=4)
    ]
    workbook.close()

    for student_number in range(3, 5):
        append_result(
            results_path,
            "Группа",
            f"Фамилия{student_number}",
            f"Имя{student_number}",
            student_number + 1,
        )

    workbook = load_workbook(results_path)
    worksheet = workbook.active
    actual_original_rows = [
        tuple(cell.value for cell in row)
        for row in worksheet.iter_rows(min_row=2, max_row=4)
    ]
    assert worksheet.max_row == 6
    assert actual_original_rows == original_rows
    assert worksheet.column_dimensions["E"].width == 22
    assert worksheet["E2"].number_format == "DD.MM.YYYY HH:MM"
    workbook.close()


def test_existing_journal_is_formatted_without_changing_records(
    tmp_path: Path,
) -> None:
    results_path = tmp_path / "results.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(RESULT_HEADERS)
    timestamp = datetime(2026, 10, 27, 19, 30)
    worksheet.append(
        ["Группа", "Иванов", "Иван", 12, timestamp, "нет"]
    )
    worksheet.column_dimensions["E"].width = 10
    workbook.save(results_path)
    workbook.close()

    assert find_first_recorded_ticket(
        results_path, "Группа", "Иванов", "Иван"
    ) == 12

    workbook = load_workbook(results_path)
    worksheet = workbook.active
    assert tuple(cell.value for cell in worksheet[2]) == (
        "Группа",
        "Иванов",
        "Иван",
        12,
        timestamp,
        "нет",
    )
    assert worksheet["E2"].number_format == "DD.MM.YYYY HH:MM"
    assert worksheet.column_dimensions["E"].width == 22
    assert worksheet.column_dimensions["A"].width >= len("Группа") + 2
    workbook.close()


def test_locked_journal_retry_keeps_assignment_without_duplicate(
    tmp_path: Path,
    valid_ticket_data: dict[int, list[str]],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    tickets_path = tmp_path / "tickets.docx"
    results_path = tmp_path / "results.xlsx"
    create_document(tickets_path, valid_ticket_data)
    choices: list[list[int]] = []

    def choose_ticket(ticket_numbers: list[int]) -> int:
        choices.append(ticket_numbers)
        return ticket_numbers[0]

    monkeypatch.setattr(ticket_service.random, "choice", choose_ticket)
    real_append_result = ticket_service.append_result
    save_attempts = 0

    def fail_first_save(*args: object, **kwargs: object) -> None:
        nonlocal save_attempts
        save_attempts += 1
        if save_attempts == 1:
            raise PermissionError("Workbook is locked")
        real_append_result(*args, **kwargs)

    monkeypatch.setattr(ticket_service, "append_result", fail_first_save)
    service = TicketService(tickets_path, results_path)

    with pytest.raises(TicketSaveError):
        service.generate_ticket("Группа", "Иванов", "Иван")

    workbook = load_workbook(results_path)
    assert workbook.active.max_row == 1
    workbook.close()

    saved_assignment = service.retry_save("Группа", "Иванов", "Иван")

    assert choices == [list(valid_ticket_data)]
    assert saved_assignment.ticket_number == choices[0][0]
    workbook = load_workbook(results_path)
    assert workbook.active.max_row == 2
    assert workbook.active["D2"].value == saved_assignment.ticket_number
    workbook.close()


def test_missing_source_files_are_reported_without_crashing(
    tmp_path: Path,
) -> None:
    assert get_group_names(tmp_path / "students.xlsx") == []
    with pytest.raises(TicketParseError, match="Не удалось прочитать"):
        parse_tickets(tmp_path / "tickets.docx")


def test_original_ticket_is_reused_from_earliest_record(
    tmp_path: Path,
    valid_ticket_data: dict[int, list[str]],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    tickets_path = tmp_path / "tickets.docx"
    results_path = tmp_path / "results.xlsx"
    create_document(tickets_path, valid_ticket_data)
    append_result(results_path, "Группа", "Иванов", "Иван", 11)
    append_result(results_path, "Группа", "Иванов", "Иван", 2)
    monkeypatch.setattr(
        ticket_service.random,
        "choice",
        lambda _numbers: pytest.fail("Existing assignment must be reused"),
    )

    assignment = TicketService(tickets_path, results_path).generate_ticket(
        "Группа", "Иванов", "Иван"
    )

    assert assignment.ticket_number == 11
    assert assignment.is_repeat is True
    workbook = load_workbook(results_path)
    assert workbook.active.max_row == 4
    assert workbook.active["D4"].value == 11
    workbook.close()
