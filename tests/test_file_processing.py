from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from excel_handler import (
    RESULT_HEADERS,
    append_result,
    find_first_recorded_ticket,
    get_group_names,
    get_students,
)
from ticket_parser import TicketParseError, parse_tickets


def test_student_workbook_groups_and_records(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Группа 1"
    worksheet.append(["Фамилия", "Имя"])
    worksheet.append(["  Иванов ", " Иван "])
    worksheet.append([None, None])
    worksheet.append(["Петров", None])
    workbook.create_sheet("Группа 2")
    file_path = tmp_path / "students.xlsx"
    workbook.save(file_path)
    workbook.close()

    assert get_group_names(file_path) == ["Группа 1", "Группа 2"]
    assert get_students(file_path, "Группа 1") == [
        {"last_name": "Иванов", "first_name": "Иван"}
    ]
    assert get_students(file_path, "Группа 2") == []
    assert get_students(tmp_path / "missing.xlsx", "Группа 1") == []

    corrupted_path = tmp_path / "corrupted.xlsx"
    corrupted_path.write_bytes(b"not an Excel workbook")
    assert get_group_names(corrupted_path) == []


def test_results_journal_appends_and_marks_repeated_records(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "results.xlsx"

    assert find_first_recorded_ticket(file_path, "Группа 1", "Иванов", "Иван") is None
    append_result(file_path, "Группа 1", "Иванов", "Иван", 8)
    assert find_first_recorded_ticket(file_path, "Группа 1", "Иванов", "Иван") == 8
    append_result(file_path, "Группа 1", "Иванов", "Иван", 3)

    workbook = load_workbook(file_path)
    worksheet = workbook.active
    assert tuple(cell.value for cell in worksheet[1]) == RESULT_HEADERS
    assert worksheet.max_row == 3
    assert worksheet["F2"].value == "нет"
    assert worksheet["F3"].value == "да"
    assert worksheet["D2"].value == 8
    assert worksheet["D3"].value == 3
    workbook.close()

    incompatible_path = tmp_path / "existing.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "Existing journal data"
    workbook.save(incompatible_path)
    workbook.close()
    with pytest.raises(ValueError):
        append_result(incompatible_path, "Группа 1", "Иванов", "Иван", 5)
    workbook = load_workbook(incompatible_path)
    assert workbook.active["A1"].value == "Existing journal data"
    workbook.close()


def test_ticket_parser_extracts_valid_tickets_and_skips_malformed(
    tmp_path: Path,
) -> None:
    document = Document()
    document.add_paragraph("Билет 7")
    document.add_paragraph("1. Первый вопрос")
    document.add_paragraph("2) Второй вопрос")
    document.add_paragraph("3. Третий вопрос")
    document.add_paragraph("Билет 20")
    document.add_paragraph("1. Только один вопрос")
    document.add_paragraph("Билет 31")
    document.add_paragraph("1. Первый вопрос")
    document.add_paragraph("2. Второй вопрос")
    document.add_paragraph("3. Третий вопрос")
    document.add_paragraph("4. Лишний вопрос")
    document.add_paragraph("Билет xyz")
    file_path = tmp_path / "tickets.docx"
    document.save(file_path)

    assert parse_tickets(file_path) == {
        7: ["Первый вопрос", "Второй вопрос", "Третий вопрос"]
    }


def test_ticket_parser_reports_missing_or_invalid_document(
    tmp_path: Path,
) -> None:
    with pytest.raises(TicketParseError):
        parse_tickets(tmp_path / "missing.docx")

    corrupted_path = tmp_path / "corrupted.docx"
    corrupted_path.write_bytes(b"not a Word document")
    with pytest.raises(TicketParseError):
        parse_tickets(corrupted_path)

    document = Document()
    document.add_paragraph("Билет 4")
    document.add_paragraph("1. Один вопрос")
    file_path = tmp_path / "invalid.docx"
    document.save(file_path)

    with pytest.raises(TicketParseError):
        parse_tickets(file_path)
