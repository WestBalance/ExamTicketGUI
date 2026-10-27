from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

import ticket_service
from excel_handler import append_result
from ticket_service import (
    TicketAssignmentError,
    TicketSaveError,
    TicketService,
)


def create_ticket_document(file_path: Path, tickets: dict[int, list[str]]) -> None:
    document = Document()
    for ticket_number, questions in tickets.items():
        document.add_paragraph(f"Билет {ticket_number}")
        for question_number, question in enumerate(questions, start=1):
            document.add_paragraph(f"{question_number}. {question}")
    document.save(file_path)


def test_student_keeps_first_ticket_across_service_instances(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    tickets_path = tmp_path / "tickets.docx"
    results_path = tmp_path / "results.xlsx"
    create_ticket_document(
        tickets_path,
        {4: ["Вопрос A", "Вопрос B", "Вопрос C"], 19: ["D", "E", "F"]},
    )
    monkeypatch.setattr(ticket_service.random, "choice", lambda numbers: 19)

    first_service = TicketService(tickets_path, results_path)
    first_request = first_service.generate_ticket("Группа", "Иванов", "Иван")
    second_request = TicketService(tickets_path, results_path).generate_ticket(
        "Группа", "Иванов", "Иван"
    )

    assert first_request.ticket_number == 19
    assert first_request.is_repeat is False
    assert second_request.ticket_number == 19
    assert second_request.is_repeat is True
    workbook = load_workbook(results_path)
    assert workbook.active.max_row == 3
    workbook.close()


def test_retry_saves_same_assignment_without_choosing_again(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    tickets_path = tmp_path / "tickets.docx"
    results_path = tmp_path / "results.xlsx"
    create_ticket_document(
        tickets_path, {4: ["Вопрос A", "Вопрос B", "Вопрос C"]}
    )
    choices: list[list[int]] = []

    def choose_ticket(numbers: list[int]) -> int:
        choices.append(numbers)
        return numbers[0]

    monkeypatch.setattr(ticket_service.random, "choice", choose_ticket)
    real_append_result = ticket_service.append_result
    save_attempts = 0

    def fail_once(*args: object, **kwargs: object) -> None:
        nonlocal save_attempts
        save_attempts += 1
        if save_attempts == 1:
            raise PermissionError("locked")
        real_append_result(*args, **kwargs)

    monkeypatch.setattr(ticket_service, "append_result", fail_once)
    service = TicketService(tickets_path, results_path)

    with pytest.raises(TicketSaveError, match="Билет 4 назначен, но не сохранён"):
        service.generate_ticket("Группа", "Иванов", "Иван")

    saved = service.retry_save("Группа", "Иванов", "Иван")
    assert saved.ticket_number == 4
    assert saved.is_repeat is False
    assert choices == [[4]]
    assert save_attempts == 2
    workbook = load_workbook(results_path)
    assert workbook.active.max_row == 2
    workbook.close()


def test_missing_original_ticket_does_not_append_or_reassign(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    tickets_path = tmp_path / "tickets.docx"
    results_path = tmp_path / "results.xlsx"
    create_ticket_document(
        tickets_path, {4: ["Вопрос A", "Вопрос B", "Вопрос C"]}
    )
    append_result(results_path, "Группа", "Иванов", "Иван", 99)
    monkeypatch.setattr(
        ticket_service.random,
        "choice",
        lambda numbers: pytest.fail("A replacement ticket must not be chosen"),
    )

    with pytest.raises(TicketAssignmentError, match="Назначить другой билет нельзя"):
        TicketService(tickets_path, results_path).generate_ticket(
            "Группа", "Иванов", "Иван"
        )

    workbook = load_workbook(results_path)
    assert workbook.active.max_row == 2
    assert workbook.active["D2"].value == 99
    workbook.close()
