import logging
import os
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.etree.ElementTree import ParseError

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.utils import get_column_letter


logger = logging.getLogger(__name__)
RESULT_HEADERS = (
    "Группа",
    "Фамилия",
    "Имя",
    "Номер билета",
    "Дата и время",
    "Повтор (да/нет)",
)


def _read_error(error: Exception, file_path: Path) -> None:
    logger.error("Не удалось прочитать файл Excel %s: %s", file_path, error)


def get_group_names(file_path: str | Path) -> list[str]:
    """Return worksheet names from the student workbook."""
    path = Path(file_path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (
        OSError,
        InvalidFileException,
        zipfile.BadZipFile,
        ParseError,
        ValueError,
    ) as error:
        _read_error(error, path)
        return []

    try:
        return list(workbook.sheetnames)
    except (
        OSError,
        InvalidFileException,
        zipfile.BadZipFile,
        ParseError,
        ValueError,
    ) as error:
        _read_error(error, path)
        return []
    finally:
        workbook.close()


def get_students(
    file_path: str | Path, group_name: str
) -> list[dict[str, str]]:
    """Return valid students from a group worksheet."""
    path = Path(file_path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (
        OSError,
        InvalidFileException,
        zipfile.BadZipFile,
        ParseError,
        ValueError,
    ) as error:
        _read_error(error, path)
        return []

    try:
        if group_name not in workbook.sheetnames:
            logger.warning("Группа %s не найдена в файле %s", group_name, path)
            return []

        worksheet = workbook[group_name]
        students = []
        for row_number, (last_name, first_name) in enumerate(
            worksheet.iter_rows(min_row=2, max_col=2, values_only=True),
            start=2,
        ):
            last_name = last_name.strip() if isinstance(last_name, str) else ""
            first_name = (
                first_name.strip() if isinstance(first_name, str) else ""
            )
            if not last_name and not first_name:
                continue
            if not last_name or not first_name:
                logger.warning(
                    "Пропущена запись студента без фамилии или имени "
                    "в группе %s, строка %s",
                    group_name,
                    row_number,
                )
                continue
            students.append(
                {"last_name": last_name, "first_name": first_name}
            )
        return students
    except (
        OSError,
        InvalidFileException,
        zipfile.BadZipFile,
        ParseError,
        ValueError,
    ) as error:
        _read_error(error, path)
        return []
    finally:
        workbook.close()


def _validate_results_workbook(workbook: Any, file_path: Path) -> None:
    worksheet = workbook.active
    headers = tuple(
        worksheet.cell(row=1, column=column).value
        for column in range(1, len(RESULT_HEADERS) + 1)
    )
    if headers != RESULT_HEADERS or worksheet.max_column > len(RESULT_HEADERS):
        raise ValueError(
            f"Файл журнала {file_path} содержит неизвестный формат. "
            "Существующие данные не изменены."
        )


def _format_results_worksheet(worksheet: Any) -> None:
    for row_number in range(2, worksheet.max_row + 1):
        worksheet.cell(row=row_number, column=5).number_format = (
            "DD.MM.YYYY HH:MM"
        )

    for column in range(1, len(RESULT_HEADERS) + 1):
        if column == 5:
            worksheet.column_dimensions[get_column_letter(column)].width = 22
            continue

        values = (
            worksheet.cell(row=row_number, column=column).value
            for row_number in range(1, worksheet.max_row + 1)
        )
        longest_value = max(
            (
                max(
                    (len(line) for line in str(value).splitlines()),
                    default=0,
                )
                for value in values
                if value is not None
            ),
            default=0,
        )
        worksheet.column_dimensions[get_column_letter(column)].width = (
            longest_value + 2
        )


def _save_atomically(workbook: Workbook, file_path: Path) -> None:
    file_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=file_path.parent,
            prefix=f".{file_path.stem}_",
            suffix=file_path.suffix or ".xlsx",
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        workbook.save(temporary_path)
        os.replace(temporary_path, file_path)
    except PermissionError as error:
        logger.error(
            "Нет доступа к файлу %s. Закройте его в Excel и повторите попытку.",
            file_path,
        )
        raise PermissionError(
            f"Файл {file_path} занят или недоступен. "
            "Закройте его в Excel и повторите попытку."
        ) from error
    finally:
        if temporary_path is not None:
            try:
                temporary_path.unlink()
            except FileNotFoundError:
                pass


def create_results_workbook(file_path: str | Path) -> None:
    """Create or format the results journal without discarding existing data."""
    path = Path(file_path)
    if path.exists():
        try:
            workbook = load_workbook(path)
        except PermissionError as error:
            logger.error("Нет доступа к файлу журнала %s: %s", path, error)
            raise PermissionError(
                f"Нет доступа к файлу журнала {path}. "
                "Закройте его в Excel и повторите попытку."
            ) from error
        except (
            OSError,
            InvalidFileException,
            zipfile.BadZipFile,
            ParseError,
            ValueError,
        ) as error:
            logger.error("Не удалось прочитать файл журнала %s: %s", path, error)
            raise ValueError(
                f"Не удалось прочитать файл журнала {path}. "
                "Существующие данные не изменены."
            ) from error
        try:
            _validate_results_workbook(workbook, path)
            _format_results_worksheet(workbook.active)
            _save_atomically(workbook, path)
        finally:
            workbook.close()
        return

    workbook = Workbook()
    try:
        worksheet = workbook.active
        worksheet.append(RESULT_HEADERS)
        _format_results_worksheet(worksheet)
        _save_atomically(workbook, path)
    finally:
        workbook.close()


def _load_results_workbook(file_path: Path) -> Workbook:
    create_results_workbook(file_path)
    try:
        workbook = load_workbook(file_path)
    except PermissionError as error:
        logger.error("Нет доступа к файлу журнала %s: %s", file_path, error)
        raise PermissionError(
            f"Нет доступа к файлу журнала {file_path}. "
            "Закройте его в Excel и повторите попытку."
        ) from error
    except (
        OSError,
        InvalidFileException,
        zipfile.BadZipFile,
        ParseError,
        ValueError,
    ) as error:
        logger.error("Не удалось прочитать файл журнала %s: %s", file_path, error)
        raise ValueError(
            f"Не удалось прочитать файл журнала {file_path}. "
            "Существующие данные не изменены."
        ) from error

    try:
        _validate_results_workbook(workbook, file_path)
    except ValueError:
        workbook.close()
        raise
    return workbook


def find_first_recorded_ticket(
    file_path: str | Path,
    group_name: str,
    last_name: str,
    first_name: str,
) -> int | str | None:
    """Return the earliest ticket number recorded for the student, if any."""
    path = Path(file_path)
    workbook = _load_results_workbook(path)
    try:
        worksheet = workbook.active
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            if (
                row[0] == group_name
                and row[1] == last_name
                and row[2] == first_name
            ):
                ticket_number = row[3]
                if isinstance(ticket_number, (int, str)):
                    return ticket_number
                return str(ticket_number) if ticket_number is not None else None
        return None
    finally:
        workbook.close()


def append_result(
    file_path: str | Path,
    group_name: str,
    last_name: str,
    first_name: str,
    ticket_number: int | str,
    requested_at: datetime | None = None,
) -> None:
    """Append a generation record and save the complete journal atomically."""
    path = Path(file_path)
    workbook = _load_results_workbook(path)
    try:
        worksheet = workbook.active
        is_repeat = any(
            row[0] == group_name
            and row[1] == last_name
            and row[2] == first_name
            for row in worksheet.iter_rows(min_row=2, values_only=True)
        )
        worksheet.append(
            [
                group_name,
                last_name,
                first_name,
                ticket_number,
                requested_at or datetime.now(),
                "да" if is_repeat else "нет",
            ]
        )
        _format_results_worksheet(worksheet)
        _save_atomically(workbook, path)
    finally:
        workbook.close()
