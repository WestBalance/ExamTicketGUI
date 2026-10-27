import os
import random
import msvcrt
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


FILE_PATH = Path(__file__).resolve().parent / "journal.xlsx"
HEADERS = ["Last name", "First name", "Номер билета", "Дата и время"]


def read_input(prompt):
    while True:
        print(prompt, end="", flush=True)
        value = ""

        while True:
            key = msvcrt.getwch()

            # ESC — выход из приложения
            if key == "\x1b":
                print("\nПрограмма завершена.")
                return None

            # Enter — завершение ввода
            elif key == "\r":
                print()
                value = value.strip()

                if value:
                    return value

                print("Поле не может быть пустым.")
                break

            # Backspace — удаление символа
            elif key == "\x08":
                if value:
                    value = value[:-1]
                    print("\b \b", end="", flush=True)

            # Игнорирование специальных клавиш
            elif key in ("\x00", "\xe0"):
                msvcrt.getwch()

            # Обычный ввод символов
            elif key.isprintable():
                value += key
                print(key, end="", flush=True)


def save_student(last_name, first_name, ticket):

    while True:
        temp_path = None

        try:
            # Открываем существующий файл или создаём новый
            if FILE_PATH.exists():
                workbook = load_workbook(FILE_PATH)
                sheet = workbook.active
            else:
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "Journal"
                sheet.append(HEADERS)

            # Добавляем новую запись
            current_time = datetime.now().strftime("%d.%m.%Y %H:%M:%S")

            sheet.append([
                last_name,
                first_name,
                ticket,
                current_time
            ])

            # Сначала сохраняем во временный файл
            with tempfile.NamedTemporaryFile(
                suffix=".xlsx",
                dir=FILE_PATH.parent,
                delete=False
            ) as temp_file:
                temp_path = Path(temp_file.name)

            workbook.save(temp_path)
            workbook.close()

            # Заменяем журнал готовым файлом
            os.replace(temp_path, FILE_PATH)

            print("Запись сохранена в journal.xlsx.\n")
            return True

        except (PermissionError, OSError) as error:
            print("\nНе удалось сохранить запись.")
            print("Возможно, journal.xlsx открыт в Excel.")
            print(f"Причина: {error}")
            print("Закройте файл и повторите попытку.")

            while True:
                choice = read_input(
                    "Введите R для повтора или ESC для выхода: "
                )

                if choice is None:
                    return False

                if choice.lower() == "r":
                    break

        finally:
            # Удаляем временный файл, если он остался
            if temp_path is not None and temp_path.exists():
                try:
                    temp_path.unlink()
                except OSError:
                    pass


def main():
    print("Генератор экзаменационных билетов")
    print("Для выхода нажмите ESC.\n")

    while True:
        # Ввод фамилии
        last_name = read_input("Last name: ")

        if last_name is None:
            break

        # Ввод имени
        first_name = read_input("First name: ")

        if first_name is None:
            break

        # Генерация случайного билета
        ticket = random.randint(1, 20)

        # Сохранение в Excel
        if save_student(last_name, first_name, ticket):
            print(f"Билет № {ticket}.\n")
        else:
            break


if __name__ == "__main__":
    main()